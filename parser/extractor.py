# [file name]: extractor.py
from googleapiclient.http import MediaIoBaseDownload
from google_services import get_drive_service, get_drive_and_sheets_services, SPREADSHEET_ID
import pandas as pd
from openpyxl import load_workbook
from pathlib import Path
import time
import logging
import asyncio
from parser.preprocess import add_year, clean, normalize_rooms, normalize_time, english_to_russian_lookalike, normalize_date_ranges, remove_invalid_dates, extract_first_date, remove_spaces_between_initials, normalize_subgroup
from parser.rooms_hints import parse_rooms_with_hints
import zipfile
from xml.etree import ElementTree as ET

logger = logging.getLogger(__name__)

pd.set_option('display.max_rows', None)
pd.set_option('display.max_columns', None)
pd.set_option('display.width', None)

_XLSX_NS = {
    'main': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
    'rel':  'http://schemas.openxmlformats.org/package/2006/relationships',
}
_R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'

def delete_old_file(latest_path, old_path, max_time=0):

    if not latest_path.exists():
        return

    file_time = latest_path.stat().st_mtime
    current_time = time.time()

    if current_time - file_time > max_time:
        if old_path.exists():
            old_path.unlink()

        latest_path.rename(old_path)


class DataExtractor:
    def __init__(self):
        self.file_id = SPREADSHEET_ID

    async def get_services(self):
        # build() синхронный и медленный (парсинг discovery-документа) — уносим в поток
        return await asyncio.to_thread(get_drive_and_sheets_services)

    async def get_sheets_metadata(self):
        _, sheets_service = await self.get_services()

        try:
            spreadsheet = sheets_service.spreadsheets().get(
                spreadsheetId=self.file_id
            ).execute()
            
            sheets_metadata = {}
            for sheet in spreadsheet.get('sheets', []):
                sheet_props = sheet['properties']
                sheets_metadata[sheet_props['title']] = {
                    'sheetId': sheet_props['sheetId'],
                    'title': sheet_props['title'],
                    'index': sheet_props['index'],
                    'gid': sheet_props['sheetId']
                }
            
            return sheets_metadata
            
        except Exception as e:
            print(f"Ошибка при получении метаданных листов: {e}")
            return {}


    def _sync_download(self, excel_path: Path) -> None:
        """Синхронный download Excel из Google Drive. Вызывать через asyncio.to_thread."""
        service = get_drive_service()

        request = service.files().export_media(
            fileId=SPREADSHEET_ID,
            mimeType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        excel_path.parent.mkdir(parents=True, exist_ok=True)
        with open(excel_path, "wb") as f:
            downloader = MediaIoBaseDownload(f, request)
            done = False
            while not done:
                status, done = downloader.next_chunk()
                logger.info(f"Download progress: {int(status.progress() * 100)}%")

    async def download_file(self, excel_path):
        await asyncio.to_thread(self._sync_download, excel_path)

    async def extract(self, file_path):
        sheets_metadata = await self.get_sheets_metadata()

        # Порядок листов берём у pandas — он совпадает с порядком в книге.
        all_sheet_names = pd.ExcelFile(file_path).sheet_names

        # Открываем книгу в read_only, чтобы дёшево узнать, какие листы скрыты.
        wb = load_workbook(file_path, read_only=True, data_only=True)
        try:
            hidden_sheets = {
                ws.title
                for ws in wb.worksheets
                if ws.sheet_state != "visible"     # 'hidden' и 'veryHidden'
            }
        finally:
            wb.close()

        skipped = [s for s in all_sheet_names if s in hidden_sheets]
        if skipped:
            logger.info(f"Пропускаю скрытые листы: {skipped}")

        visible_sheets = [s for s in all_sheet_names if s not in hidden_sheets]

        groups_info = {}

        # Первый лист — сводный ("ГРУППЫ"), не парсим.
        for sheet_name in visible_sheets[1:]:
            df = fill_merged_cells_safe(file_path, sheet_name=sheet_name)
            hyperlinks = load_hyperlinks(file_path, sheet_name)
            sheet_gid = None
            if sheet_name in sheets_metadata:
                sheet_gid = sheets_metadata[sheet_name]['gid']
            # if not sheet_name == "1214":
            #     continue

            group_info = extraction(df, sheet_gid, hyperlinks)
            if not group_info:
                continue

            # Не затираем уже собранные группы сводными листами.
            for grp, data in group_info.items():
                if grp in groups_info and groups_info[grp].get("events"):
                    continue
                groups_info[grp] = data

        return groups_info


def clean_group_name(group_name):
    bracket_index = group_name.find(' (')
    comma_index = group_name.find(', ')+2
    if bracket_index != -1:
        return group_name[comma_index:bracket_index].strip()
    return group_name.strip()

def extraction(df, sheet_id, hyperlinks: dict | None = None):
    if hyperlinks is None:
        hyperlinks = {}
    # print(df)
    # Указываем стартовые клетки
    start_cells = []
    for i, row in df.iterrows():
        # row = set(row)
        for j, cell in enumerate(row):
            if "Начало" in str(cell).strip():
                start_cells.append((i, j))
    # Указываем конечную высоту поиска
    end_row = None
    for i, row in df.iterrows():
        row = set(row)
        for j, cell in enumerate(row):
            if "Декан" in str(cell).strip():
                end_row = i
                break
        if end_row:
            break
    # Указываем конечные клетки
    end_cells = []
    for i, row in df.iterrows():
        # row = set(row)
        for j, cell in enumerate(row):
            if "форма обучения /" in str(cell).strip().lower():
                end_cells.append((end_row,j))
    # Ищем заголовки
    headers = []
    for i, row in df.iterrows():
        row_set = set()
        for j, cell in enumerate(row):
            if "семестр" in str(cell).strip().lower() and not (cell in row_set):
                headers.append((i,j))
                row_set.add(cell)

    group_info = {}

    min_count = min(len(start_cells), len(end_cells), len(headers))


    if min_count == 0:
        return group_info

    # Проходим по всем блокам по каждой клетке внутри блока
    for i in range(min_count):

        start_cell = start_cells[i]
        end_cell = end_cells[i]
        header_cell = headers[i]

        try:
            # Разбираем заголовок
            header = df.iloc[header_cell[0], header_cell[1]]

            group_name = ""
            semester = None
            additional_info = ""
            session = ""
            vacation = ""
            session_end_date = None

            if "семестр" in str(header).lower():
                lines = [line.strip() for line in str(header).split("\n") if line.strip()]

                if len(lines) >= 2:
                    group_name = lines[0]
                    group_name = clean_group_name(group_name)
                    semester = lines[1]

                    try:
                        session_index = next(i for i, line in enumerate(lines) if "Экзаменационная сессия" in line)
                        vacation_index = next(i for i, line in enumerate(lines) if "Каникулы" in line)

                        additional_info = "\n".join(lines[2:session_index])
                        session = lines[session_index]
                        session_end_date = extract_first_date(session)
                        vacation = lines[vacation_index]
                    except StopIteration:
                        print("Не найдены сессия или каникулы в заголовке")
                else:
                    print(f"Недостаточно строк в заголовке: {len(lines)}")
        except Exception as e:
            print(f"Ошибка при обработке заголовка: {e}")
            continue

        if not group_name:
            print("Не удалось определить номер группы, пропускаем")
            continue

        try:
            subset = df.iloc[start_cell[0]:end_cell[0]+1, start_cell[1]:end_cell[1]+1]
        except Exception as e:
            print(f"Ошибка при создании subset: {e}")
            continue

        # times = []
        # subjects = []
        # teachers = []
        # rooms = []

        events = []

        # Разбираем строчки на пары
        for row in range(1, subset.shape[0]-1):
            if row >= subset.shape[0]:
                break

            abs_row = start_cell[0] + row
            # time_col = start_cell[1]
            # subject_col = start_cell[1] + 1
            # room_col = end_cell[1] - 1

            try:

                # extra
                time_val_exception = df.iloc[abs_row, start_cells[0][1]]
                room_val_exception = df.iloc[abs_row, end_cells[-1][1]]

                time_val = subset.iloc[row, 0]
                room_val = subset.iloc[row, -1]
                subject_val = subset.iloc[row, 1:-1]

                if pd.isna(time_val) or str(time_val).strip() == "":
                    continue
                if pd.isna(room_val) or str(room_val).strip() == "" or room_val is None:
                    room_val = ""

                subject_val = set(subject_val)
                subject_val = " ".join([str(x).strip() for x in subject_val if pd.notna(x)]).strip()
                if not subject_val:
                    continue

                if (time_val in subject_val) or (room_val in subject_val) and (room_val != "" and subject_val != ""):
                    if time_val in subject_val:
                        time_val = time_val_exception
                    if room_val in subject_val:
                        room_val = room_val_exception

                row_urls: list[str] = []
                for c in range(start_cell[1], end_cell[1] + 3):
                    for u in (hyperlinks.get((abs_row, c)) or []):
                        if u not in row_urls:
                            row_urls.append(u)

                room_raw = str(room_val)
                logger.info(f"[EXTRACT] group={group_name} row={abs_row} row_urls={row_urls} "
                            f"raw_repr={room_raw!r}")
                rooms_hints = parse_rooms_with_hints(room_raw, urls=row_urls)

                if rooms_hints:
                    room_val = [r["room"] for r in rooms_hints]
                else:
                    room_val = normalize_rooms(
                        room_raw, url=row_urls[0] if row_urls else None
                    )

                subject_val = clean(text=subject_val)
                subject_val = english_to_russian_lookalike(text=subject_val)
                subject_val = normalize_date_ranges(text=subject_val, default_end_date=session_end_date)
                subject_val = remove_invalid_dates(text=subject_val)
                subject_val = add_year(text=subject_val)

                subject_val = remove_spaces_between_initials(text=subject_val)
                subject_val = normalize_subgroup(text=subject_val)
                # subject_val = remove_academic_titles(text=subject_val)
                # subject_val = clean(text=subject_val)

                # Добавляем день недели
                day_of_week = ""
                df.iloc[:, 0] = df.iloc[:, 0].ffill()
                left_col_val = df.iloc[abs_row, 0]
                if pd.notna(left_col_val) and str(left_col_val).strip():
                    day_of_week = str(left_col_val).strip().upper()


                    time_vals = time_val.split("\n")

                    merged_times = []
                    i = 0
                    while i < len(time_vals):
                        cur = time_vals[i].strip()
                        if cur.endswith(("-", "–", "—")) and i + 1 < len(time_vals):
                            nxt = time_vals[i + 1].strip()
                            merged_times.append(cur + nxt)
                            i += 2
                        else:
                            merged_times.append(cur)
                            i += 1

                    for time_value in merged_times:
                        time_val = normalize_time(time_value)
                        # НОВЫЙ ФОРМАТ: dict вместо строки-обёртки
                        events.append({
                            "day_of_week": day_of_week,
                            "time_start": time_val,
                            "rooms": ", ".join(room_val),
                            "rooms_hints": rooms_hints,
                            "subject": subject_val,
                        })

            except Exception as e:
                print(f"Ошибка при обработке строки {row} {group_name}: {e}")

        # print(events)
        if events:
            group_info[group_name] = {
                "events": events,
            }

    return group_info

def fill_merged_cells_safe(file_path, sheet_name):
    """
    Reads an Excel sheet and fills only the cells that belong to merged ranges.
    Genuine NaN values outside merged ranges are left untouched.
    """
    # 1. Read the sheet with pandas (merged cells appear as NaN except top-left)
    df = pd.read_excel(file_path, sheet_name=sheet_name, header=None)
    df = df.astype(object)

    # 2. Load the same sheet with openpyxl to get merged range definitions
    wb = load_workbook(file_path, data_only=True)
    if isinstance(sheet_name, int):
        ws = wb.worksheets[sheet_name]
    else:
        ws = wb[sheet_name]

    for merged_range in ws.merged_cells.ranges:
        top_left_value = ws.cell(merged_range.min_row, merged_range.min_col).value
        start_row = merged_range.min_row - 1
        end_row = merged_range.max_row - 1
        start_col = merged_range.min_col - 1
        end_col = merged_range.max_col - 1

        if start_row < 0 or end_row >= len(df) or start_col < 0 or end_col >= len(df.columns):
            continue

        for i in range(start_row, end_row + 1):
            for j in range(start_col, end_col + 1):
                df.iat[i, j] = top_left_value
    return df

def _col_idx(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + (ord(ch) - ord('A') + 1)
    return n - 1


def _parse_ref(ref: str):
    """'G17' → (row0, col0)."""
    import re
    m = re.match(r'^([A-Z]+)(\d+)$', ref)
    if not m:
        return None
    return int(m.group(2)) - 1, _col_idx(m.group(1))


def _parse_range(ref: str):
    """'A1:B2' → (min_r, min_c, max_r, max_c), 0-based."""
    import re
    m = re.match(r'^([A-Z]+)(\d+):([A-Z]+)(\d+)$', ref)
    if not m:
        return None
    return (int(m.group(2)) - 1, _col_idx(m.group(1)),
            int(m.group(4)) - 1, _col_idx(m.group(3)))


def _find_sheet_path(zf: zipfile.ZipFile, sheet_name):
    wb   = ET.fromstring(zf.read('xl/workbook.xml'))
    rels = ET.fromstring(zf.read('xl/_rels/workbook.xml.rels'))

    rid_to_target = {
        r.get('Id'): r.get('Target')
        for r in rels.findall('rel:Relationship', _XLSX_NS)
    }

    sheets = wb.findall('.//main:sheet', _XLSX_NS)

    if isinstance(sheet_name, int):
        sh = sheets[sheet_name] if 0 <= sheet_name < len(sheets) else None
    else:
        sh = next((s for s in sheets if s.get('name') == sheet_name), None)

    if sh is None:
        return None

    rid = sh.get(f'{{{_R_NS}}}id')
    target = rid_to_target.get(rid)
    if not target:
        return None

    target = target.lstrip('/')
    if not target.startswith('xl/'):
        target = f'xl/{target}'
    return target

def load_hyperlinks(file_path, sheet_name) -> dict[tuple[int, int], list[str]]:
    """
    {(row, col): [url, url, ...]} — координаты 0-based (как в pandas).

    Читает XML напрямую, потому что openpyxl через cell.hyperlink отдаёт
    ровно одну ссылку на ячейку. При rich-text-гиперссылках (несколько
    ссылок в одной ячейке) остальные теряются.

    Гиперссылки в объединённых ячейках распространяются на всю область.
    """
    links: dict[tuple[int, int], list[str]] = {}

    with zipfile.ZipFile(file_path) as zf:
        sheet_path = _find_sheet_path(zf, sheet_name)
        if sheet_path is None:
            return links

        d, f = sheet_path.rsplit('/', 1)
        rels_path = f'{d}/_rels/{f}.rels'

        rid_to_url: dict[str, str] = {}
        try:
            rels = ET.fromstring(zf.read(rels_path))
            for rel in rels.findall('rel:Relationship', _XLSX_NS):
                rid  = rel.get('Id')
                url  = rel.get('Target')
                typ  = rel.get('Type') or ''
                if rid and url and typ.endswith('/hyperlink'):
                    rid_to_url[rid] = url
        except KeyError:
            pass

        sheet = ET.fromstring(zf.read(sheet_path))

        # --- прямые <hyperlink ref=".." r:id=".."/> ---
        for h in sheet.findall('.//main:hyperlink', _XLSX_NS):
            ref = h.get('ref')
            if not ref:
                continue
            rid = h.get(f'{{{_R_NS}}}id')
            url = rid_to_url.get(rid) if rid else None
            if not url or not url.lower().startswith(('http://', 'https://')):
                continue

            pos = _parse_ref(ref)
            if pos is not None:
                links.setdefault(pos, [])
                if url not in links[pos]:
                    links[pos].append(url)
                continue

            rng = _parse_range(ref)
            if rng is None:
                continue
            min_r, min_c, max_r, max_c = rng
            for r in range(min_r, max_r + 1):
                for c in range(min_c, max_c + 1):
                    links.setdefault((r, c), [])
                    if url not in links[(r, c)]:
                        links[(r, c)].append(url)

        # --- распространение по объединённым ячейкам ---
        merged: list[tuple[int, int, int, int]] = []
        for mc in sheet.findall('.//main:mergeCell', _XLSX_NS):
            rng = _parse_range(mc.get('ref') or '')
            if rng:
                merged.append(rng)

        for min_r, min_c, max_r, max_c in merged:
            src = (min_r, min_c)
            if src not in links:
                continue
            for r in range(min_r, max_r + 1):
                for c in range(min_c, max_c + 1):
                    if (r, c) == src:
                        continue
                    links.setdefault((r, c), [])
                    for u in links[src]:
                        if u not in links[(r, c)]:
                            links[(r, c)].append(u)

    return links
