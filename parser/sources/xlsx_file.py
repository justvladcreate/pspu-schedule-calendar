from __future__ import annotations

import logging
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pandas as pd
from openpyxl import load_workbook

from parser.sources.base import ScheduleSource, SheetData

logger = logging.getLogger(__name__)

_XLSX_NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "rel":  "http://schemas.openxmlformats.org/package/2006/relationships",
}
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

def _col_idx(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + (ord(ch) - ord("A") + 1)
    return n - 1

def _parse_ref(ref: str):
    m = re.match(r"^([A-Z]+)(\d+)$", ref)
    if not m:
        return None
    return int(m.group(2)) - 1, _col_idx(m.group(1))

def _parse_range(ref: str):
    m = re.match(r"^([A-Z]+)(\d+):([A-Z]+)(\d+)$", ref)
    if not m:
        return None
    return (
        int(m.group(2)) - 1, _col_idx(m.group(1)),
        int(m.group(4)) - 1, _col_idx(m.group(3)),
    )

def _sheet_paths_from_zip(path: str | Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as zf:
        wb   = ET.fromstring(zf.read("xl/workbook.xml"))
        rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))

    rid_to_target = {
        r.get("Id"): r.get("Target")
        for r in rels.findall("rel:Relationship", _XLSX_NS)
    }

    result: dict[str, str] = {}
    for sh in wb.findall(".//main:sheet", _XLSX_NS):
        name = sh.get("name")
        rid = sh.get(f"{{{_R_NS}}}id")
        target = rid_to_target.get(rid)
        if not name or not target:
            continue
        target = target.lstrip("/")
        if not target.startswith("xl/"):
            target = f"xl/{target}"
        result[name] = target
    return result

def _is_gray_hex(hex_str: str | None) -> bool:
    if not hex_str:
        return False
    hex_str = hex_str.upper()
    if hex_str.startswith('FF'):
        hex_str = hex_str[2:]
    if len(hex_str) != 6:
        return False
    try:
        r = int(hex_str[0:2], 16)
        g = int(hex_str[2:4], 16)
        b = int(hex_str[4:6], 16)
    except ValueError:
        return False
    mx, mn = max(r, g, b), min(r, g, b)
    if mx - mn > 20:
        return False
    return 30 < mx < 230

def _read_sheet_df(excel_file: pd.ExcelFile, worksheet) -> tuple[pd.DataFrame, dict[tuple[int, int], list[str]]]:
    df = excel_file.parse(worksheet.title, header=None)
    df = df.astype(object)

    for merged_range in worksheet.merged_cells.ranges:
        top_left_value = worksheet.cell(merged_range.min_row, merged_range.min_col).value
        start_row = merged_range.min_row - 1
        end_row   = merged_range.max_row - 1
        start_col = merged_range.min_col - 1
        end_col   = merged_range.max_col - 1

        if start_row < 0 or end_row >= len(df) or start_col < 0 or end_col >= len(df.columns):
            continue

        for i in range(start_row, end_row + 1):
            for j in range(start_col, end_col + 1):
                df.iat[i, j] = top_left_value

    colors = {}
    for row in worksheet.iter_rows():
        for cell in row:
            color = cell.font.color
            if color and color.type == 'rgb' and _is_gray_hex(color.rgb):
                if cell.value:
                    colors[(cell.row - 1, cell.column - 1)] = [str(cell.value)]

    return df, colors

def _load_hyperlinks(zf: zipfile.ZipFile, sheet_path: str) -> dict[tuple[int, int], list[str]]:
    links: dict[tuple[int, int], list[str]] = {}
    if not sheet_path:
        return links

    d, f = sheet_path.rsplit("/", 1)
    rels_path = f"{d}/_rels/{f}.rels"

    rid_to_url: dict[str, str] = {}
    try:
        rels = ET.fromstring(zf.read(rels_path))
        for rel in rels.findall("rel:Relationship", _XLSX_NS):
            rid = rel.get("Id")
            url = rel.get("Target")
            typ = rel.get("Type") or ""
            if rid and url and typ.endswith("/hyperlink"):
                rid_to_url[rid] = url
    except KeyError:
        pass

    sheet = ET.fromstring(zf.read(sheet_path))

    for h in sheet.findall(".//main:hyperlink", _XLSX_NS):
        ref = h.get("ref")
        if not ref:
            continue
        rid = h.get(f"{{{_R_NS}}}id")
        url = rid_to_url.get(rid) if rid else None
        if not url or not url.lower().startswith(("http://", "https://")):
            continue

        pos = _parse_ref(ref)
        if pos is not None:
            links.setdefault(pos, [])
            if url not in links[pos]:
                links[pos].append(url)
            continue

        rng = _parse_range(ref)
        if rng is None:
            continue
        min_r, min_c, max_r, max_c = rng
        for r in range(min_r, max_r + 1):
            for c in range(min_c, max_c + 1):
                links.setdefault((r, c), [])
                if url not in links[(r, c)]:
                    links[(r, c)].append(url)

    for mc in sheet.findall(".//main:mergeCell", _XLSX_NS):
        rng = _parse_range(mc.get("ref") or "")
        if not rng:
            continue
        min_r, min_c, max_r, max_c = rng
        src = (min_r, min_c)
        if src not in links:
            continue
        for r in range(min_r, max_r + 1):
            for c in range(min_c, max_c + 1):
                if (r, c) == src:
                    continue
                links.setdefault((r, c), [])
                for u in links[src]:
                    if u not in links[(r, c)]:
                        links[(r, c)].append(u)

    return links


class XlsxFileSource(ScheduleSource):
    def __init__(self, path: str | Path):
        self.path = Path(path)

    async def load(self) -> list[SheetData]:
        if not self.path.exists():
            raise FileNotFoundError(f"Нет файла: {self.path}")

        wb = load_workbook(self.path, data_only=True)
        excel_file = pd.ExcelFile(self.path)

        try:
            hidden = {
                ws.title for ws in wb.worksheets if ws.sheet_state != "visible"
            }
            sheet_paths = _sheet_paths_from_zip(self.path)

            result: list[SheetData] = []
            with zipfile.ZipFile(self.path) as zf:
                for name in wb.sheetnames:
                    df, colors = _read_sheet_df(excel_file, wb[name])
                    links = _load_hyperlinks(zf, sheet_paths.get(name))
                    result.append(SheetData(
                        name=name,
                        df=df,
                        links=links,
                        colors=colors,
                        hidden=(name in hidden),
                        sheet_id=None,
                    ))
            return result
        finally:
            wb.close()
            excel_file.close()