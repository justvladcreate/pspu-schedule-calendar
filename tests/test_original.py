import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from parser.core.original import publish_original, render_sheet_html


def _workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "131"

    sheet["A1"] = "РАСПИСАНИЕ УЧЕБНЫХ ЗАНЯТИЙ"
    sheet["A1"].font = Font(name="Times New Roman", size=14, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.merge_cells("A1:F1")

    sheet["A3"] = "Дисциплина\nвторая строка"
    sheet["A3"].border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    sheet["B3"] = "практика"
    sheet["B3"].font = Font(color="0000FF", italic=True, bold=True)
    sheet["B3"].fill = PatternFill("solid", fgColor="FFFF00")

    sheet.column_dimensions["A"].width = 20
    sheet.row_dimensions[1].height = 30
    return workbook


def test_render_sheet_html_keeps_merges_and_styles():
    html = render_sheet_html(_workbook().active)

    assert "colspan=\"6\"" in html
    assert "font-family:Times New Roman" in html
    assert "font-weight:700" in html
    assert "text-align:center" in html
    assert "border-top:1px solid #000000" in html
    assert "color:#0000FF" in html
    assert "background:#FFFF00" in html
    assert "font-style:italic" in html
    assert "Дисциплина<br>вторая строка" in html
    assert "height:40px" in html  # 30pt → 40px


def test_render_skips_hidden_sheets(tmp_path):
    workbook = _workbook()
    hidden = workbook.create_sheet("служебный")
    hidden["A1"] = "не показывать"
    hidden.sheet_state = "hidden"

    source = tmp_path / "source.xlsx"
    workbook.save(source)

    entries = publish_original(tmp_path / "web", "01", source)

    assert [entry["name"] for entry in entries] == ["131"]
    folder = tmp_path / "web" / "faculties" / "01"
    index = json.loads((folder / "sheets.json").read_text(encoding="utf-8"))
    assert index["sheets"][0]["file"].startswith("sheets/")
    assert (folder / index["sheets"][0]["file"]).exists()


def test_publish_original_without_file(tmp_path):
    assert publish_original(tmp_path, "01", tmp_path / "missing.xlsx") is None
